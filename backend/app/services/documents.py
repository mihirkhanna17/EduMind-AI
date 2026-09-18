"""Document ingestion (PDF via PyMuPDF, slides via python-pptx) and retrieval.

Retrieval is Postgres full-text search for MVP — the pgvector `embedding`
column is populated when an embedding provider is configured (Poe has no
embeddings endpoint; this is the documented seam to add one).
"""

from __future__ import annotations

import io

import fitz  # PyMuPDF
from pptx import Presentation
from sqlalchemy import select, text
from sqlalchemy.ext.asyncio import AsyncSession

from app.models import Document, DocumentChunk

CHUNK_SIZE = 1200
CHUNK_OVERLAP = 150


def extract_text(filename: str, data: bytes) -> tuple[str, str]:
    """Returns (source_type, full_text)."""
    lower = filename.lower()
    if lower.endswith(".pdf"):
        with fitz.open(stream=data, filetype="pdf") as doc:
            return "pdf", "\n\n".join(page.get_text() for page in doc)
    if lower.endswith(".pptx"):
        prs = Presentation(io.BytesIO(data))
        slides = []
        for slide in prs.slides:
            parts = [
                shape.text_frame.text
                for shape in slide.shapes
                if shape.has_text_frame and shape.text_frame.text.strip()
            ]
            slides.append("\n".join(parts))
        return "pptx", "\n\n".join(slides)
    raise ValueError(f"Unsupported file type: {filename} (pdf/pptx only)")


def chunk_text(full_text: str) -> list[str]:
    chunks = []
    start = 0
    while start < len(full_text):
        end = start + CHUNK_SIZE
        chunk = full_text[start:end].strip()
        if chunk:
            chunks.append(chunk)
        start = end - CHUNK_OVERLAP
    return chunks


async def ingest_document(
    db: AsyncSession, user_id: int, filename: str, data: bytes
) -> Document:
    source_type, full_text = extract_text(filename, data)
    doc = Document(user_id=user_id, filename=filename, source_type=source_type)
    db.add(doc)
    await db.flush()
    for chunk in chunk_text(full_text):
        db.add(DocumentChunk(document_id=doc.id, content=chunk))
    await db.commit()
    return doc


async def search_chunks(db: AsyncSession, user_id: int, query: str, limit: int = 4) -> list[str]:
    """Full-text search over the user's uploaded material."""
    rows = (
        await db.execute(
            text(
                """
                SELECT dc.content
                FROM document_chunks dc
                JOIN documents d ON d.id = dc.document_id
                WHERE d.user_id = :user_id
                  AND to_tsvector('english', dc.content) @@ plainto_tsquery('english', :q)
                ORDER BY ts_rank(to_tsvector('english', dc.content),
                                 plainto_tsquery('english', :q)) DESC
                LIMIT :limit
                """
            ),
            {"user_id": user_id, "q": query, "limit": limit},
        )
    ).all()
    return [r[0] for r in rows]
